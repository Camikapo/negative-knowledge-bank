"""
Экспорт результата генерации гипотез в Word-документ (.docx).

Использует python-docx. Возвращает байты, готовые для скачивания или сохранения.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# --- Цветовая палитра, согласованная с интерфейсом ------------------------
ACCENT = RGBColor(0x2B, 0x4C, 0x3F)      # тёмно-зелёный акцент
MUTED = RGBColor(0x6B, 0x6B, 0x68)        # приглушённый серый
TEXT = RGBColor(0x1F, 0x1F, 0x1F)         # почти чёрный
WARM = RGBColor(0xB0, 0x8D, 0x57)         # тёплый охристый для меток "неудача"


def _shade_cell(cell, hex_color: str) -> None:
    """Заливка ячейки таблицы (python-docx не поддерживает это из коробки)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _add_label(paragraph, text: str) -> None:
    """Маленькая подпись капсом акцентным цветом — как в интерфейсе."""
    run = paragraph.add_run(text.upper())
    run.font.size = Pt(8)
    run.font.color.rgb = ACCENT
    run.font.bold = True


def build_report_docx(result: dict[str, Any]) -> bytes:
    """Строит DOCX-отчёт из результата generate_hypotheses и возвращает bytes.

    Структура: титульная секция → блок цели → таблица учтённых неудач →
    развёрнутые карточки гипотез → колонтитул.
    """
    doc = Document()

    # Уменьшаем поля страницы — воздух остаётся, но не пустует
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # Базовый шрифт документа
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = TEXT

    # ---------------- Титульный блок ---------------------------------------
    p = doc.add_paragraph()
    _add_label(p, "Negative Knowledge Bank")

    h = doc.add_paragraph()
    run = h.add_run("Отчёт по гипотезам исследования")
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = TEXT

    meta = doc.add_paragraph()
    ts = datetime.now().strftime("%d.%m.%Y, %H:%M")
    run = meta.add_run(f"Сгенерировано {ts}")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    # ---------------- Цель и ограничения -----------------------------------
    doc.add_paragraph()
    p = doc.add_paragraph()
    _add_label(p, "Цель исследования")
    p = doc.add_paragraph(result.get("goal", "—"))
    p.paragraph_format.space_after = Pt(6)

    if result.get("constraints"):
        p = doc.add_paragraph()
        _add_label(p, "Ограничения")
        doc.add_paragraph(result["constraints"])

    # ---------------- Учтённые прошлые неудачи -----------------------------
    doc.add_paragraph()
    h = doc.add_paragraph()
    run = h.add_run("Учтённые прошлые неудачи")
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = TEXT

    failures = result.get("relevant_failures", [])
    if failures:
        table = doc.add_table(rows=1 + len(failures), cols=3)
        table.autofit = False

        # Ширины колонок
        widths = [Cm(2.2), Cm(11.0), Cm(2.5)]
        for row in table.rows:
            for i, cell in enumerate(row.cells):
                cell.width = widths[i]

        # Заголовок
        hdr = table.rows[0].cells
        for i, title in enumerate(("ID", "Название", "Релевантность")):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            run = p.add_run(title)
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = MUTED
            _shade_cell(hdr[i], "F0EEE9")

        # Строки
        for idx, rf in enumerate(failures, start=1):
            row = table.rows[idx].cells
            row[0].text = rf.get("id", "—")
            row[1].text = rf.get("title", "—")
            row[2].text = f"{rf.get('score', 0):.2f}"
            for c in row:
                for para in c.paragraphs:
                    for r in para.runs:
                        r.font.size = Pt(10)
    else:
        p = doc.add_paragraph("Похожих неудач в банке не найдено — гипотезы построены с нуля.")
        p.runs[0].font.color.rgb = MUTED

    # ---------------- Гипотезы ---------------------------------------------
    doc.add_paragraph()
    h = doc.add_paragraph()
    run = h.add_run("Сгенерированные гипотезы")
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = TEXT

    for i, hyp in enumerate(result.get("hypotheses", []), start=1):
        # Заголовок карточки
        p = doc.add_paragraph()
        num = p.add_run(f"{i}.  ")
        num.font.size = Pt(12)
        num.font.bold = True
        num.font.color.rgb = ACCENT
        st = p.add_run(hyp.get("statement", "—"))
        st.font.size = Pt(12)
        st.font.bold = True
        st.font.color.rgb = TEXT

        # Пары «подпись → значение»
        pairs = [
            ("Механизм влияния", hyp.get("mechanism")),
            ("Опирается на прошлую неудачу", hyp.get("based_on_failure")),
            ("Чем отличается от того, что не сработало", hyp.get("why_different")),
            ("Новизна", hyp.get("novelty")),
            ("Риски", hyp.get("risks")),
            ("Ожидаемая ценность", hyp.get("expected_value")),
            ("Как проверить в лаборатории", hyp.get("verification")),
        ]
        for label, value in pairs:
            if not value:
                continue
            lp = doc.add_paragraph()
            lp.paragraph_format.space_before = Pt(4)
            lp.paragraph_format.space_after = Pt(1)
            _add_label(lp, label)
            vp = doc.add_paragraph(str(value))
            vp.paragraph_format.left_indent = Cm(0.3)
            vp.paragraph_format.space_after = Pt(2)

        # Разделитель между гипотезами
        sep = doc.add_paragraph()
        run = sep.add_run("—" * 40)
        run.font.color.rgb = RGBColor(0xE5, 0xE3, 0xDD)
        run.font.size = Pt(8)

    # ---------------- Футер -------------------------------------------------
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = footer.add_run(
        "Negative Knowledge Bank · автоматически сгенерировано на базе Yandex AI Studio"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    # ---------------- В байты -----------------------------------------------
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
